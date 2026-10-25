from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=r'C:\QLDA\DatVeXe\output\bai-tra-loi\Nguyễn Xuân Phú - 2380601708 - Quản lý nguồn nhân lực.docx'
d=Document()
s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
for n in ['Normal','Title','Heading 1']:
 st=d.styles[n];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(12)
 st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.0
 st.paragraph_format.widow_control=True
st=d.styles['Heading 1'];st.font.bold=True;st.font.size=Pt(12);st.paragraph_format.space_before=Pt(7)
def p(t):
 a=d.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if ': ' in t:
  x,y=t.split(': ',1);a.add_run(x+': ').bold=True;a.add_run(y)
 else:a.add_run(t)
 return a
p('Họ tên: Nguyễn Xuân Phú    |    MSSV: 2380601708')
a=d.add_paragraph('Quản lý nguồn nhân lực cho dự án','Title');a.runs[0].font.size=Pt(15);a.runs[0].bold=True;a.paragraph_format.space_after=Pt(6)
d.add_paragraph('Câu 1 Các bước quản lý nguồn nhân lực và tầm quan trọng','Heading 1')
for t in [
'1. Lập kế hoạch nhân lực: Xác định công việc, số lượng người, năng lực, vai trò, thời điểm tham gia và chi phí. Bước này giúp bố trí đúng người, tránh thiếu hoặc dư nhân lực và bảo đảm tiến độ.',
'2. Huy động và thành lập nhóm: Tuyển chọn, điều chuyển hoặc thuê nhân sự phù hợp; thống nhất nhiệm vụ và thời gian làm việc. Đây là cơ sở để dự án có đủ người và kỹ năng cần thiết khi triển khai.',
'3. Phát triển nhóm: Đào tạo, hướng dẫn, xây dựng tinh thần hợp tác và động viên thành viên. Bước này giúp nâng cao năng lực, tăng sự gắn kết và cải thiện chất lượng công việc.',
'4. Quản lý và đánh giá nhóm: Theo dõi khối lượng, kết quả, phản hồi và xử lý xung đột; điều chỉnh phân công khi cần. Việc này giúp phát hiện khó khăn sớm, duy trì hiệu suất và tránh quá tải.',
'5. Kết thúc và giải phóng nhân lực: Đánh giá đóng góp, bàn giao công việc, ghi nhận bài học và bố trí lại nhân sự. Bước này bảo đảm chuyển giao đầy đủ và sử dụng kinh nghiệm cho dự án sau.'
]:p(t)
d.add_paragraph('Câu 2 Vai trò của ma trận phân công trách nhiệm RAM','Heading 1')
p('RAM liên kết các công việc với cá nhân hoặc bộ phận phụ trách. Ma trận làm rõ ai thực hiện, ai quyết định và ai cần phối hợp, qua đó hạn chế chồng chéo, bỏ sót nhiệm vụ và hỗ trợ kiểm soát khối lượng công việc. Khi nhân sự thay đổi, RAM giúp xác định trách nhiệm cần bàn giao; ma trận không thay thế lịch tiến độ hoặc kế hoạch năng lực.')
p('Ví dụ: Trong dự án đặt vé xe, RAM có thể thể hiện theo dạng RACI: R là người thực hiện; A là người chịu trách nhiệm cuối cùng và phê duyệt; C là người được tham vấn; I là người được thông báo. Mỗi công việc nên có một A rõ ràng.')
t=d.add_table(rows=1, cols=5);t.autofit=False
widths=[6,2.7,2.7,2.8,2.8]
for c,w in zip(t.columns,widths):c.width=Cm(w)
for c,txt in zip(t.rows[0].cells,['Công việc','Quản lý dự án','Lập trình viên','Kiểm thử viên','Khách hàng']):c.text=txt
for vals in [('Xây dựng chức năng đặt vé','A','R','C','I'),('Kiểm thử chức năng đặt vé','A','C','R','I')]:
 for c,txt in zip(t.add_row().cells,vals):c.text=txt
for ri,row in enumerate(t.rows):
 for ci,c in enumerate(row.cells):
  tcPr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  tcPr.append(borders)
  if ri==0:
   e=OxmlElement('w:shd');e.set(qn('w:fill'),'E7E6E6');tcPr.append(e)
  for a in c.paragraphs:
   a.paragraph_format.space_after=Pt(3);a.paragraph_format.space_before=Pt(3)
   a.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
   for r in a.runs:r.font.size=Pt(10);r.bold=ri==0
p('Nhờ ma trận trên, lập trình viên biết mình triển khai chức năng, kiểm thử viên biết mình kiểm tra; quản lý dự án chịu trách nhiệm cuối cùng cho cả hai công việc.')
d.add_paragraph('Câu 3 Tầm quan trọng và yếu tố của kế hoạch tổ chức','Heading 1')
p('Kế hoạch tổ chức xác định cơ cấu nhóm, quan hệ báo cáo và quyền hạn của từng vị trí. Kế hoạch rõ ràng giúp ra quyết định nhanh, phối hợp thống nhất, phân bổ nhân lực phù hợp và giải quyết xung đột trách nhiệm. Nếu thiếu kế hoạch, thành viên có thể nhận chỉ đạo trái ngược hoặc không biết báo cáo cho ai.')
p('Các yếu tố cần xem xét: Mục tiêu, phạm vi và tiến độ dự án; khối lượng công việc và kỹ năng cần có; ngân sách, số lượng và mức độ sẵn sàng của nhân sự; cơ cấu tổ chức và quyền điều phối của quản lý dự án; vai trò, trách nhiệm và quyền phê duyệt; kênh trao đổi, quan hệ báo cáo; địa điểm, múi giờ và văn hóa làm việc; nhu cầu đào tạo, động viên, nhân sự dự phòng và yêu cầu lao động liên quan.')
d.core_properties.author='Nguyễn Xuân Phú';d.core_properties.title='Quản lý nguồn nhân lực cho dự án'
d.save(out)
print(out)
print('Words:',sum(len(p.text.split()) for p in d.paragraphs))
