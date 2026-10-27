from docx import Document
from pathlib import Path
import re,json
b=Path(__file__).parent;d=Document(b/'Báo cáo quản lý dự án DatVeXe - hoàn tất.docx')
ids=[re.match(r'TASK-(\d+)',r.cells[0].text).group(1) for t in d.tables for r in t.rows if re.match(r'TASK-\d+ [EST]$',r.cells[0].text)]
assert len(ids)==119 and len(set(ids))==119
fields=[e.text for e in d._element.xpath('.//w:instrText')]
assert sum('TOC ' in f for f in fields)==3
assert sum('SEQ ' in f for f in fields)==45
assert len(d.inline_shapes)==11
text='\n'.join(p.text for p in d.paragraphs)
assert 'Cập nhật danh mục' not in text and 'Error!' not in text
assert len(d.sections)==3
(b/'verified.json').write_text(json.dumps({'WBS':119,'captions':45,'TOC_and_lists':3,'body_images':11,'sections':3,'status':'PASS'},indent=2),encoding='utf-8')
print('PASS: 119 công việc, 45 caption tự động, 3 danh mục, 11 hình, 3 section')
