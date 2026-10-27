from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from copy import deepcopy
from pathlib import Path
from io import BytesIO
import re, hashlib, json

base=Path(__file__).parent
d=Document(base/'original.docx'); ref=Document(base/'reference.docx')
ps=list(d.paragraphs); tables=list(d.tables)
replacements={
122:'Báo cáo tổng hợp quá trình quản lý dự án DatVeXe từ ngày 07/09/2026 đến khi kết thúc Sprint 6 ngày 27/10/2026. Theo kịch bản hoàn tất của báo cáo, toàn bộ 119 công việc đã được hoàn thành; giai đoạn 28/10–01/11 được dành làm thời gian dự trữ và hoàn thiện hồ sơ. Các số liệu chi phí và phát hiện lỗi ở từng kỳ được giữ theo hồ sơ theo dõi, không đồng nhất trạng thái hoàn thành với việc không còn lỗi.',
128:'Nhóm đã sử dụng Jira để phân công và theo dõi công việc, GitHub để quản lý mã nguồn, Excel để lập WBS, CPM, EMV và Pareto. Dự án được tổ chức thành sáu Sprint từ ngày 14/09 đến 27/10/2026. Trước đó, nhóm thực hiện khảo sát, xác định yêu cầu và chuẩn bị backlog trong giai đoạn 07–13/09. Sau Sprint 6, nhóm chuyển sang tổng hợp sản phẩm và hồ sơ bàn giao.',
138:'Trong quá trình thực hiện, nhóm thống nhất các thay đổi về phạm vi, thời gian và phân công thông qua trao đổi trực tiếp. Nguyễn Xuân Phú tổng hợp quyết định và cập nhật các tài liệu liên quan để các thành viên sử dụng cùng một kế hoạch.',
150:'WBS của dự án gồm 10 Epic, 27 Story và 82 Subtask, tổng cộng 119 công việc. Đến cuối Sprint 6, các công việc được xem là đã hoàn thành theo kịch bản báo cáo. Epic và Story là cấp tổng hợp; công sức và chi phí chỉ được cộng một lần từ các công việc lá. Phụ lục A giữ đầy đủ mã, tên và quan hệ phân cấp của danh mục WBS.',
154:'Nhóm đã hoàn thành các hạng mục bàn giao gồm mã nguồn Frontend và Backend, cấu trúc cơ sở dữ liệu, tài liệu nghiệp vụ, giao diện và bộ hồ sơ quản lý dự án. Việc kết thúc phạm vi được xác định ở cuối Sprint 6, với 119/119 công việc hoàn thành. Các yêu cầu về luồng đặt vé, phân quyền, tính nhất quán dữ liệu và khả năng tái lập môi trường tiếp tục là cơ sở đối chiếu chất lượng sản phẩm khi bàn giao.',
158:'3.2 Tiến độ thực hiện sáu Sprint',
161:'Lịch sáu Sprint được giữ theo Jira và được dùng làm đường cơ sở để tổng hợp tiến độ. Mốc kết thúc Sprint 6 là ngày 27/10/2026; toàn bộ phạm vi công việc được xem là hoàn tất tại mốc này. Giai đoạn dự trữ 28/10–01/11 không được tính thành một Sprint mới. Ngày hoàn thành của từng công việc không được sửa ngược chỉ để trùng ngày kết thúc Sprint.',
170:'Mạng Subtask có 210 đường găng cùng thời lượng 41 ngày công. Vì nhiều nhánh cùng găng, việc tăng tốc một nhánh riêng lẻ chưa chắc rút ngắn toàn bộ dự án. Trong quá trình theo dõi, nhóm tập trung vào các công việc tích hợp và phụ thuộc giữa các phần để hạn chế ảnh hưởng dây chuyền đến thời điểm bàn giao.',
184:'4.3 Tổng hợp chi phí và cơ sở quyết toán',
186:'Phần chi phí đã có số liệu được giữ theo hồ sơ Jira đến ngày 25/10/2026. Giả định hoàn tất 119 công việc cho phép tổng hợp tiến độ và giá trị công việc hoàn thành tại ngày 27/10, nhưng không tự xác định thêm số ngày công hoặc khoản chi sau ngày 25/10. Vì vậy, tổng chi phí cuối kỳ chỉ được quyết toán sau khi đối chiếu phần ngày công còn lại và các khoản hỗ trợ thực chi.',
191:'Nhóm lập báo cáo EMV lũy kế sau các Sprint 2, 3, 4 và 6, tương ứng ngày 26/09, 05/10, 10/10 và 27/10/2026. Mỗi kỳ bao quát 119 công việc, tổng hợp từ 82 Subtask lên Story và Epic. Với giả định toàn bộ công việc hoàn thành vào cuối Sprint 6, PV và EV của kỳ cuối cùng bằng ngân sách công việc 103 triệu đồng; SPI bằng 1 và SV bằng 0. AC cuối kỳ chưa được xác định từ dữ liệu ngày công đầy đủ, nên CPI và CV cuối kỳ chưa được kết luận. Các kỳ trước giữ số liệu lịch sử để phản ánh diễn biến thực hiện.',
196:'Các biểu đồ Gantt và báo cáo EMV trong hồ sơ quản lý được tổ chức theo sáu Sprint. Đường thẳng đứng màu đỏ thể hiện ngày báo cáo để đối chiếu phần công việc kế hoạch với kết quả tại cùng mốc. Đối với COCOMO, nhóm chưa có KLOC được xác nhận nên không bổ sung một ước lượng thiếu căn cứ vào phần tổng kết.',
199:'Nhóm xây dựng tiêu chí chất lượng theo luồng nghiệp vụ, tính nhất quán dữ liệu, phân quyền và khả năng triển khai. Trong kịch bản hoàn tất, các hạng mục phát triển và kiểm thử thuộc WBS đã hoàn thành. Kết quả đó cần được đọc cùng với các phát hiện của từng kỳ rà soát; trạng thái hoàn thành công việc không đồng nghĩa phần mềm không còn hạn chế.',
203:'Nhóm đã tổ chức công việc theo sáu Sprint. Trước mỗi Sprint, các thành viên thống nhất mục tiêu và phạm vi; trong Sprint, nhóm cập nhật tiến độ và vướng mắc; sau Sprint, nhóm kiểm tra demo, trao đổi nguyên nhân phát sinh và cách xử lý. Đến cuối Sprint 6, không còn công việc thuộc phạm vi WBS ở trạng thái chưa hoàn thành trong kịch bản tổng kết này. Điểm đánh giá cá nhân được xem xét riêng theo tiêu chí và minh chứng của từng thành viên.',
205:'Pareto được tổ chức riêng cho từng Sprint, với mốc rà soát lần lượt là 20/09, 26/09, 05/10, 10/10, 20/10 và 27/10/2026. Hồ sơ hiện có là kết quả kiểm tra hồi cứu trên các snapshot lịch sử của nhánh Develops. Các phát hiện có thể lặp lại giữa các kỳ và không được cộng thành tổng số lỗi mới độc lập. Việc hoàn thành công việc trong báo cáo không làm thay đổi số lỗi đã ghi nhận tại các mốc trước.',
226:'Sprint 6 kết thúc ngày 27/10/2026, với phạm vi kiểm thử, sửa lỗi, tài liệu, triển khai và bàn giao đã hoàn thành theo kịch bản báo cáo. Hồ sơ nguồn chưa có kết quả rà soát mã nguồn cuối kỳ tương ứng ngày này, nên không gán số lỗi bằng 0 hoặc tạo phân bố Pareto giả định. Các biểu đồ trước đó vẫn phản ánh kết quả của đúng snapshot được kiểm tra.',
228:'Qua các kỳ rà soát, nhóm ưu tiên xử lý vấn đề phân quyền, sai dữ liệu và tính nhất quán của giao dịch trước khi xử lý các nhóm lỗi về mã nguồn và giao diện. Ngưỡng 80% hỗ trợ xác định nhóm nguyên nhân xuất hiện nhiều, nhưng mức độ nghiêm trọng vẫn là căn cứ ưu tiên chính. Kết thúc dự án, nhóm tổng hợp các phát hiện, phần xử lý và những điểm cần tiếp tục theo dõi trong hồ sơ bàn giao.',
239:'Phân công chuyên môn giúp các thành viên phối hợp trong những phần việc có quan hệ phụ thuộc. Nguyễn Xuân Phú điều phối yêu cầu và tích hợp; Nguyễn Hữu Dũng phụ trách trải nghiệm người dùng và triển khai; Lê Văn Phú phụ trách API, dữ liệu và kiểm thử. Việc đánh giá cá nhân dựa trên chất lượng công việc, tiến độ, khả năng phối hợp và minh chứng trên Jira, GitHub.',
243:'Toàn bộ 82 Subtask đã hoàn thành theo kịch bản tổng kết. Ngày công WBS phản ánh phần công sức được phân bổ cho từng công việc; quỹ lương hai tháng là cơ sở ngân sách, còn chi phí thực tế được tính theo ngày công thực hiện. Các hoạt động họp, hỗ trợ, sửa lỗi và hoàn thiện tài liệu cần được đối chiếu riêng khi đánh giá tổng tải của từng thành viên.',
281:'Sản phẩm DatVeXe gồm Frontend React/Vite, Backend Express và cơ sở dữ liệu SQL Server được truy cập qua Prisma. Nhóm đã hoàn thành các công việc phát triển, tích hợp, tài liệu và triển khai trong phạm vi WBS. Mã nguồn và các tài liệu dưới đây là thành phần hồ sơ bàn giao; hướng dẫn cấu hình và chạy môi trường được lưu trong README của dự án.',
287:'Các hình dưới đây minh họa giao diện và cấu trúc cơ sở dữ liệu của DatVeXe trong hồ sơ sản phẩm. Luồng thể hiện từ tìm chuyến, chọn ghế đến thông tin hành khách và xác nhận đặt vé. Phần thanh toán được giới hạn trong phạm vi giao diện đã thiết kế; không mở rộng kết luận sang tích hợp thanh toán ngân hàng thực tế ngoài MVP.',
307:'Kết thúc Sprint 6 ngày 27/10/2026, dự án hoàn tất 119 công việc gồm 10 Epic, 27 Story và 82 Subtask theo kịch bản tổng kết. Nhóm hoàn thành phạm vi phát triển và bộ hồ sơ bàn giao trước hạn 01/11/2026, giữ giai đoạn 28/10–01/11 làm dự trữ. Tổng ngân sách được duyệt là 104.785.000 đồng; giá trị công việc hoàn thành theo ngân sách là 103.000.000 đồng. Việc quyết toán chi phí nhân công và chi hỗ trợ được thực hiện trên số ngày công và khoản chi thực tế, không lấy ngân sách làm chi phí đã phát sinh.',
309:'Phụ lục liệt kê đầy đủ 119 công việc theo thứ tự WBS. Trong kịch bản tổng kết, toàn bộ danh mục đã hoàn thành vào cuối Sprint 6; các cột ngày dưới đây giữ lịch kế hoạch để đối chiếu, không được đổi thành ngày hoàn thành thực tế. E là Epic, S là Story và T là Subtask. Giai đoạn chuẩn bị 07–13/09 và dự trữ 28/10–01/11 không tạo thêm mã công việc ngoài danh mục này.',
331:'[1] Báo cáo mẫu Báo+cáo+DACS.doc do nhóm cung cấp, tham khảo định dạng trang bìa, đề mục, mục lục và caption.',
337:'[7] Dữ liệu lịch sử Jira và hồ sơ rà soát snapshot GitHub của DatVeXe; số liệu chi phí đã tổng hợp đến ngày 25/10/2026.',
}
for i,t in replacements.items(): ps[i].text=t
ps[191].text='Báo cáo EMV được lập lũy kế sau Sprint 2, 3, 4 và 6. Số liệu được tổng hợp một lần từ 82 Subtask lên 27 Story và 10 Epic. Với giả định hoàn thành toàn bộ công việc ngày 27/10, PV = EV = 103 triệu đồng, SPI = 1 và SV = 0. AC cuối kỳ cần được đối chiếu thêm nên chưa kết luận CPI và CV. Các kỳ trước giữ số liệu lịch sử.'
ps[228].text='Nhóm ưu tiên vấn đề phân quyền, dữ liệu và giao dịch, sau đó xử lý nhóm lỗi mã nguồn và giao diện. Ngưỡng 80% hỗ trợ chọn nhóm nguyên nhân phổ biến; mức nghiêm trọng vẫn là căn cứ ưu tiên. Các phát hiện và phần cần theo dõi được tổng hợp trong hồ sơ bàn giao.'
for i in [214,219,224]:ps[i].text='Chi tiết phát hiện, vị trí mã nguồn, phương pháp và commit được lưu trong sheet Pareto tương ứng để đối chiếu và theo dõi xử lý.'
ps[190].text+=' Trên Gantt, đường thẳng đứng màu đỏ đánh dấu ngày báo cáo. Nhóm chưa lập ước lượng COCOMO vì chưa có KLOC được xác nhận.'
for i in [195,196]:ps[i]._p.getparent().remove(ps[i]._p)
# Final-period facts permitted by the completion assumption; no invented actual days/costs.
for j,t in enumerate(['6','27/10','103,000','103,000','Chưa chốt','—','1,000']): tables[12].rows[-1].cells[j].text=t
tables[12].cell(0,0).text='Sprint'
tables[11].rows[-1].cells[2].text='3.000.000'
tables[11].rows[-1].cells[1].text='Chưa chốt'
tables[11].rows[-1].cells[3].text='Chưa quyết toán'
tables[3].rows[5].cells[0].text='Chức năng mở rộng trong WBS'
tables[9].rows[1].cells[2].text='Ba thành viên trong hai tháng'
tables[9].rows[2].cells[2].text='Ba tài khoản trong hai tháng'

# Restore genuine automatic TOC/list fields instead of hand-written entries.
def field(p,code,cached=''):
    for typ in ['begin', 'instruction', 'separate', 'text', 'end']:
        r=OxmlElement('w:r')
        if typ in ('begin','separate','end'):
            e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),typ)
        elif typ=='instruction':
            e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=' '+code+' '
        else:
            e=OxmlElement('w:t');e.text=cached
        r.append(e);p._p.append(r)
for a,b,code in [(32,82,'TOC \\o "1-3" \\h \\z \\u'),(84,105,'TOC \\h \\z \\c "Bảng"'),(107,118,'TOC \\h \\z \\c "Hình"')]:
    p=ps[a];p.clear();p.style='Normal';field(p,code,'Cập nhật danh mục')
    for i in range(a+1,b+1): ps[i]._p.getparent().remove(ps[i]._p)

# Captions follow the sample: label, chapter.sequence, colon, description.
for ti,title in [(14,'Thống kê nhóm phát hiện tại Sprint 3'),(15,'Thống kê nhóm phát hiện tại Sprint 4'),(16,'Thống kê nhóm phát hiện tại Sprint 5')]+[(25+i,'Danh mục công việc thuộc '+ps[310+i*2].text) for i in range(10)]:
    el=OxmlElement('w:p');tables[ti]._tbl.addprevious(el);cp=Paragraph(el,d._body);cp.style='Table Caption';cp.text=f'Bảng {100+ti}. {title}'
chapter=0; counters={}; oldmap={}; capids=[]
for p in d.paragraphs:
    if p.text.startswith('Phụ lục A '):chapter='A'
    m=re.match(r'Chương (\d+) ',p.text)
    if p.style.name=='Heading 1' and m: chapter=int(m.group(1));p.text=f'CHƯƠNG {chapter}. '+p.text[m.end():].upper()
    if p.style.name in ('Table Caption','Figure Caption'):
        m=re.match(r'(Bảng|Hình) (\d+)\.\s*(.*)',p.text)
        if not m: continue
        label,num,title=m.groups();key=(chapter,label);counters[key]=counters.get(key,0)+1;n=counters[key]
        oldmap[f'{label} {num}']=f'{label} {chapter}.{n}'
        p.clear();p.add_run(f'{label} {chapter}.');field(p,f'SEQ {label}'+(' \\r 1' if n==1 else ''),str(n));p.add_run(': '+title)
        capids.append((label,chapter,n,title));p.paragraph_format.keep_with_next=(label=='Bảng')
    elif p.style.name=='Normal' and not p._p.xpath('.//w:drawing') and not p._p.xpath('.//w:instrText'):
        p.text=re.sub(r'\b(Bảng|Hình) \d+\b',lambda m:oldmap.get(m.group(0),m.group(0)),p.text) if ('Bảng ' in p.text or 'Hình ' in p.text) else p.text

# Copy the first reference cover including the original school logo/table.
cover=[]
end=ref.paragraphs[31]._p
for e in ref._element.body:
    cover.append(deepcopy(e))
    if e is end: break
for e in cover:
    for el in e.iter():
        for attr,val in list(el.attrib.items()):
            if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and val in ref.part.rels:
                rel=ref.part.rels[val]
                if rel.reltype.endswith('/image'):
                    from docx.opc.part import Part
                    from docx.opc.packuri import PackURI
                    part=Part(PackURI('/word/media/ref_cover_'+str(len(d.part.rels))+'.'+str(rel.target_part.partname).split('.')[-1]),rel.target_part.content_type,rel.target_part.blob,d.part.package)
                    rid=d.part.relate_to(part,rel.reltype);el.set(attr,rid)
    for p in e.iter(qn('w:p')):
        props=p.find(qn('w:pPr'))
        if props is not None:
            for st in list(props.findall(qn('w:pStyle'))): props.remove(st)
    for t in e.iter(qn('w:t')):
        if t.text:
            t.text=t.text.replace('ĐỒ ÁN CƠ SỞ','BÁO CÁO QUẢN LÝ DỰ ÁN CNTT').replace('HỆ THỐNG LẬP KẾ HOẠCH DU LỊCH THÔNG MINH','HỆ THỐNG HỖ TRỢ ĐẶT LỊCH, ĐẶT CHỖ TRONG VÉ XE').replace('Trần Thị Vân Anh','Nguyễn Hữu Trung').replace('TP. Hồ Chí Minh, 2026','TP. Hồ Chí Minh, tháng 10 năm 2026')
# Add the third student by cloning the existing reference row paragraph.
student=deepcopy(cover[-1])
for e in cover:
    if 'Nguyễn Hữu Dũng' in ''.join(e.itertext()):
        student=deepcopy(e);break
for t in student.iter(qn('w:t')):
    if t.text:t.text=t.text.replace('Nguyễn Hữu Dũng','Lê Văn Phú').replace('2380600299','2380601701').replace('23DTHC4','')
idx=next(i for i,e in enumerate(cover) if 'Nguyễn Hữu Dũng' in ''.join(e.itertext()))
cover.insert(idx+1,student)
# No unsupported class for the added member; remove class column from all names.
for e in cover:
    for t in e.iter(qn('w:t')):
        if t.text:t.text=t.text.replace('23DTHC4','').replace('Lớp:', '')
# Remove current cover while retaining all body content/images/tables.
body=d._element.body
while body[0] is not ps[11]._p: body.remove(body[0])
for e in reversed(cover):body.insert(0,e)
for e in cover:
    for el in e.iter(qn('w:p')):
        cp=Paragraph(el,d._body);cp.paragraph_format.space_before=Pt(0);cp.paragraph_format.space_after=Pt(0);cp.paragraph_format.line_spacing=1
        if not cp.text:
            rp=el.get_or_add_pPr(); sz=OxmlElement('w:rPr'); z=OxmlElement('w:sz');z.set(qn('w:val'),'20');sz.append(z);rp.append(sz)
        for r in cp.runs:
            r.font.name='Times New Roman'
            if not cp.text:r.font.size=Pt(10)
coversect=deepcopy(ref.sections[0]._sectPr)
for e in list(coversect):
    if e.tag in (qn('w:headerReference'),qn('w:footerReference')):coversect.remove(e)
p=OxmlElement('w:p');pr=OxmlElement('w:pPr');pr.append(coversect);p.append(pr);body.insert(len(cover),p)
# Section between preliminary matter and chapters: real page restart.
frontsect=deepcopy(d.sections[-1]._sectPr)
pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),'lowerRoman');pg.set(qn('w:start'),'1');frontsect.append(pg)
p=OxmlElement('w:p');pr=OxmlElement('w:pPr');pr.append(frontsect);p.append(pr);ps[119]._p.addprevious(p)
last=d.sections[-1]._sectPr
pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),'decimal');pg.set(qn('w:start'),'1');last.append(pg)

# Typography derived from the reference's visible body, not its inconsistent imported style definitions.
for sec in d.sections:
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2.5);sec.left_margin=Cm(3);sec.right_margin=Cm(2)
for name in ['Normal','Heading 1','Heading 2','Heading 3','Front Heading','Table Caption','Figure Caption','toc 1','toc 2','toc 3']:
    if name not in d.styles:continue
    s=d.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(13);s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.3
    if name.startswith('Heading'):
        s.font.bold=True;s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(12)
    if name=='Heading 1':s.font.size=Pt(16);s.paragraph_format.page_break_before=True
    if name=='Heading 2':s.font.size=Pt(14)
    if name=='Front Heading':s.font.size=Pt(16);s.font.bold=True
    if 'Caption' in name:s.font.size=Pt(12);s.font.italic=True;s.paragraph_format.line_spacing=1.0;s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if name.startswith('toc'):s.paragraph_format.line_spacing=1.15;s.paragraph_format.space_after=Pt(3)

cover_set=set(cover)
for p in d.paragraphs:
    if p._p in cover_set:continue
    if p.style.name=='Normal' and p.text:
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent=Cm(.75)
        p.paragraph_format.line_spacing=1.15
    if p.style.name.startswith('Heading') or p.style.name=='Front Heading':
        p.paragraph_format.first_line_indent=Cm(0)
    if p.style.name=='Heading 1':p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if p.style.name=='Heading 3' and p.text.startswith('TASK-'):
        lvl=OxmlElement('w:outlineLvl');lvl.set(qn('w:val'),'9');p._p.get_or_add_pPr().append(lvl)
    if 'Caption' in p.style.name:p.paragraph_format.first_line_indent=Cm(0)
    # Remove inherited colored or undersized direct formatting except hyperlinks.
    for r in p.runs:
        if r.text:
            r.font.name='Times New Roman';r.font.size=None;r.font.color.rgb=RGBColor(0,0,0)
    if p._p.xpath('.//w:drawing'):
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)

# Tables: repeating header and no split rows; compact appendix remains readable.
for ti,t in enumerate(d.tables):
    if ti==0:continue # original reference school logo table
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit');trpr.append(no)
        if ri==0:el=OxmlElement('w:tblHeader');trpr.append(el)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                if ri==0:p.paragraph_format.keep_with_next=True
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10 if ti>=25 else 11);r.font.color.rgb=RGBColor(0,0,0)
# Fit preserved images to the slightly narrower template area.
for shape in d.inline_shapes:
    if shape.width>Cm(16):
        ratio=Cm(16)/shape.width;shape.height=int(shape.height*ratio);shape.width=Cm(16)
up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');d.settings.element.append(up)
out=base/'Báo cáo quản lý dự án DatVeXe - hoàn tất.docx';d.save(out)
(base/'artifact.md').write_text('Reference: '+str(base/'reference.docx')+'\nSHA256: '+hashlib.sha256((base/'reference.docx').read_bytes()).hexdigest()+'\nA4 portrait; margins left 3 right 2 top 2 bottom 2.5 cm. Reference school logo and first cover cloned with project, supervisor and three students substituted. One cover retained rather than duplicate covers. Body Times New Roman 13 pt; headings 16/14/13 pt, black and bold. Caption label chapter.sequence: title, automatic SEQ; separate automatic lists by label. Main report content, all tables, 11 figures and 119-work appendix retained. Cover no page number; preliminary pages Roman; body restart Arabic. User-requested completion scenario affects narrative and final EV/PV, not missing actual costs or defect evidence.\n',encoding='utf-8')
(base/'edits.json').write_text(json.dumps({'paragraph_edits':list(replacements),'captions':capids,'output':str(out)},ensure_ascii=False,indent=2),encoding='utf-8')
print(out)

